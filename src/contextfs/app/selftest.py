"""Offline end-to-end self test of the app (ADR-0004 §3, Decision 95's light acceptance check).

1. Installs a Python audit hook that **fails on any network use** (socket connect / DNS).
2. Reports the runtime: frozen or not, ONNX Runtime version and providers, model files, and the
   Smart App Control state (Evaluation mode is an open risk, Decision 95).
3. Builds a throwaway corpus and index in a temp folder, updates, searches, re-updates
   (expects a no-op), and checks the expected results.

Exit code 0 only if every check passes.
"""

from __future__ import annotations

import sys
import tempfile
import time
from datetime import datetime
from pathlib import Path

__all__ = ["run_selftest", "smart_app_control_state"]


def _block_network() -> None:
    def hook(event, args):
        if event in ("socket.connect", "socket.getaddrinfo", "socket.bind"):
            raise RuntimeError(f"network use blocked by selftest: {event} {args!r}"[:200])

    sys.addaudithook(hook)


def smart_app_control_state() -> str:
    """'off' | 'on' | 'evaluation' | 'unknown' (read-only registry query)."""
    if sys.platform != "win32":
        return "n/a"
    try:
        import winreg

        with winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, r"SYSTEM\CurrentControlSet\Control\CI\Policy") as k:
            v, _ = winreg.QueryValueEx(k, "VerifiedAndReputablePolicyState")
        return {0: "off", 1: "on", 2: "evaluation"}.get(int(v), f"unknown({v})")
    except OSError:
        return "unknown"


def loaded_dlls() -> list[str]:
    """Paths of every DLL loaded in this process (Windows only; read-only psapi query)."""
    if sys.platform != "win32":
        return []
    import ctypes
    from ctypes import wintypes

    psapi, k32 = ctypes.WinDLL("psapi"), ctypes.WinDLL("kernel32")
    k32.GetCurrentProcess.restype = wintypes.HANDLE
    proc = k32.GetCurrentProcess()
    mods = (wintypes.HMODULE * 2048)()
    need = wintypes.DWORD()
    psapi.EnumProcessModules.argtypes = [wintypes.HANDLE, ctypes.c_void_p, wintypes.DWORD,
                                         ctypes.POINTER(wintypes.DWORD)]
    psapi.EnumProcessModules(proc, mods, ctypes.sizeof(mods), ctypes.byref(need))
    psapi.GetModuleFileNameExW.argtypes = [wintypes.HANDLE, wintypes.HMODULE, wintypes.LPWSTR, wintypes.DWORD]
    out = []
    for h in mods[: need.value // ctypes.sizeof(wintypes.HMODULE)]:
        buf = ctypes.create_unicode_buffer(1024)
        psapi.GetModuleFileNameExW(proc, h, buf, 1024)
        out.append(buf.value)
    return out


def foreign_code() -> tuple[list[str], list[str]]:
    """Frozen build only: ``(bad, injected)``.

    ``bad``: Python modules from outside the app folder, or Python/dev-environment DLLs from
    outside it (a leak from .venv/.tools or another Python). ``injected``: other DLLs from outside
    the app and Windows, which other software (antivirus, audio drivers) injects into every process.
    """
    import os

    app = Path(sys.executable).resolve().parent
    windir = Path(os.environ.get("SYSTEMROOT", r"C:\Windows")).resolve()
    bad, injected = set(), set()
    for f in filter(None, (getattr(m, "__file__", None) for m in list(sys.modules.values()))):
        p = Path(f)
        # in-archive modules (PyInstaller's PYZ / bootstrap) have names, not real files on disk
        if p.is_absolute() and p.exists() and not p.resolve().is_relative_to(app):
            bad.add(str(p.resolve()))
    for f in loaded_dlls():
        p = Path(f).resolve()
        if p.is_relative_to(app) or p.is_relative_to(windir):
            continue
        low = str(p).lower()
        dev = any(s in low for s in ("site-packages", ".venv", ".tools", "python"))
        (bad if dev or p.suffix.lower() == ".pyd" else injected).add(str(p))
    return sorted(bad), sorted(injected)


def _minimal_pdf(text: str) -> bytes:
    """A one-page PDF with one line of text, written by hand (no PDF writer is bundled)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return bytes(out)


def _office_files(root: Path) -> None:
    """One .docx, .xlsx and .pdf: proves the format extractors are bundled in a frozen build."""
    import docx
    from openpyxl import Workbook

    d = docx.Document()
    d.add_paragraph("Internship offer letter from Infosys, joining in Mysuru.")
    d.save(str(root / "College" / "offer_letter.docx"))
    wb = Workbook()
    wb.active.append(["semester", "sgpa"])
    wb.active.append(["fifth", "8.7"])
    wb.save(str(root / "College" / "grades.xlsx"))
    (root / "College" / "kmeans_clustering.pdf").write_bytes(
        _minimal_pdf("K-means clustering lecture handout with elbow method"))


def run_selftest(console=None, gpu: bool = False) -> int:
    """Run all checks; print a report; return a process exit code.

    ``gpu=True`` also runs the GPU parity probe on the chosen discrete adapter (real GPU work).
    """
    say = console.print if console else print
    _block_network()
    failures = []

    def check(name: str, ok: bool, detail: str = "") -> None:
        say(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f" - {detail}" if detail else ""))
        if not ok:
            failures.append(name)

    import onnxruntime as ort

    from contextfs.engine import Engine, default_model_dir

    say("ContextFS self test")
    say(f"  frozen app: {bool(getattr(sys, 'frozen', False))}  python {sys.version.split()[0]}")
    say(f"  onnxruntime {ort.__version__} (providers: {', '.join(ort.get_available_providers())}); queries always on the CPU")
    sac = smart_app_control_state()
    say(f"  Smart App Control: {sac}" + ("  (OPEN RISK: may switch to On and block unsigned builds)"
                                         if sac == "evaluation" else ""))
    from contextfs.engine import gpu as gpumod

    found = gpumod.adapters()
    pick = gpumod.choose_adapter(found)
    say("  GPU adapters: " + (", ".join(f"[{a['device_id']}] {a['name']} ({a['dedicated_mb']} MB)"
                                         for a in found) or "none found"))
    say(f"  GPU for indexing: {pick['name'] if pick else 'none suitable, CPU'}")
    model = default_model_dir()
    check("model files present", (model / "model_fp32.onnx").exists() and (model / "tokenizer.json").exists(),
          str(model))
    as_of = datetime(2026, 9, 30, 21).astimezone()
    with tempfile.TemporaryDirectory(prefix="contextfs-selftest-") as tmp:
        root = Path(tmp) / "files"
        (root / "College").mkdir(parents=True)
        (root / "College" / "dbms_normalization.md").write_text(
            "# BCNF and 3NF\nDecomposition examples. Submission deadline 24 November 2025.", encoding="utf-8")
        (root / "Personal").mkdir()
        (root / "Personal" / "biryani_recipe.txt").write_text("Biryani with saffron and mint.", encoding="utf-8")
        (root / "Personal" / "trip_budget.csv").write_text("item,cost\ntrain,900\nhotel,2400\n", encoding="utf-8")
        _office_files(root)
        t = time.perf_counter()
        with Engine.open(Path(tmp) / "index", mode="write") as eng:
            rep = eng.update([root])
            check("first update indexes 6 files", rep.new == 6 and rep.errors == 0, f"{rep.seconds}s")
            rep2 = eng.update([root])
            check("second update is a no-op", rep2.new == rep2.changed == rep2.deleted == 0)
        with Engine.open(Path(tmp) / "index") as eng:
            r = eng.search("notes on BCNF decomposition", as_of=as_of)
            check("content query finds the right file", bool(r.results) and r.results[0].rel_path.endswith(
                "dbms_normalization.md"))
            for query, name in (("internship offer from Infosys", "offer_letter.docx"),
                                ("my sgpa spreadsheet", "grades.xlsx"),
                                ("elbow method handout", "kmeans_clustering.pdf")):
                top = eng.search(query, as_of=as_of).paths[:1]
                check(f"{name.rsplit('.', 1)[1]} text is extracted and found", top == [f"College/{name}"],
                      str(top))
            check("every result explains itself", all(x.reasons for x in r.results))
            check("gibberish returns a clear note", not eng.search("qzxv wplk", as_of=as_of).results)
        check("end-to-end time", True, f"{time.perf_counter() - t:.1f}s")
        if gpu:
            from contextfs.engine.encoder import OnnxEncoder

            enc, note = gpumod.gpu_encoder(model, OnnxEncoder(model))
            check("GPU probe (falls back to CPU safely if it fails)", True, note)
    if getattr(sys, "frozen", False):
        bad, injected = foreign_code()
        check("no code loaded from a dev environment or another Python", not bad, "; ".join(bad[:5]))
        say(f"  loaded: {len(sys.modules)} Python modules, {len(loaded_dlls())} DLLs")
        for p in injected:
            say(f"  note: DLL injected by other software (not ours): {p}")
    say("RESULT: " + ("ALL PASS" if not failures else f"FAILED ({', '.join(failures)})"))
    return 0 if not failures else 1
