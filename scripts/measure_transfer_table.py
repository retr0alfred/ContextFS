"""Measure the transfer table T: what real tools do to every time field (ADR-0006 E1).

Fixture files with known file-system times (created < modified < accessed, all in the past) and
known in-file times (Office core.xml, PDF Info, EXIF, zip entries) are pushed through every
transfer tool available on this machine. For each (tool, format) the script records what happened
to birth, mtime and atime, whether each in-file time survived, and whether a Mark-of-the-Web
``Zone.Identifier`` stream appeared.

T is frozen (sha256 recorded) BEFORE any estimator rule is written. Rows that need a human or
hardware this script cannot drive (OneDrive, USB FAT32/exFAT, browser download, Word Save As,
7-Zip) are listed as ``pending_user`` and filled in by hand, as a later version of T, still before
the estimator rules.

Usage:
    python scripts/measure_transfer_table.py [--work DIR] [--out data/transfer_table]
"""

from __future__ import annotations

import argparse
import hashlib
import json
import platform
import shutil
import subprocess
import sys
import time
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from contextfs.datagen.filetimes import (  # noqa: E402
    get_file_times,
    read_infile_times,
    set_file_times,
)

# Known source times (local wall clock). created < modified < accessed, all well in the past.
SRC_CREATED = datetime(2025, 3, 14, 8, 0, 0)
SRC_MODIFIED = datetime(2025, 3, 14, 10, 0, 0)
SRC_ACCESSED = datetime(2025, 3, 15, 9, 0, 0)
INFILE_CREATED = "2023-01-10T09:00:00Z"
INFILE_MODIFIED = "2025-03-14T04:30:00Z"  # = 10:00 IST
EXIF_ORIGINAL = "2025:03:14 10:00:00"
ZIP_ENTRY = (2025, 3, 14, 10, 0, 0)
NOW_TOLERANCE = timedelta(minutes=5)

FORMATS = ["txt", "docx", "pdf", "jpg", "zip"]
PENDING_USER = [
    (
        "explorer_drag_copy_gui",
        "Explorer drag-and-drop copy (GUI) — the Shell COM row is the scripted proxy",
    ),
    (
        "onedrive_roundtrip",
        "Upload to OneDrive, delete local, re-download (Files-On-Demand hydrate)",
    ),
    ("usb_fat32_roundtrip", "Copy to a FAT32 USB stick and back"),
    ("usb_exfat_roundtrip", "Copy to an exFAT USB stick and back"),
    (
        "browser_download",
        "Download the file from a web page/mail attachment in a browser (Mark-of-the-Web)",
    ),
    ("word_save_as", "Open in Microsoft Word and Save As a new name (rewrites core.xml)"),
    ("sevenzip_roundtrip", "7-Zip archive and extract"),
    ("whatsapp_telegram_desktop", "Receive the file through a messenger desktop client"),
]


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


def make_fixtures(src: Path) -> dict[str, Path]:
    src.mkdir(parents=True, exist_ok=True)
    files: dict[str, Path] = {}

    p = src / "notes.txt"
    p.write_text("transfer-table fixture\n", encoding="utf-8")
    files["txt"] = p

    from docx import Document

    p = src / "report.docx"
    doc = Document()
    doc.add_paragraph("transfer-table fixture")
    cp = doc.core_properties
    cp.created = datetime.fromisoformat(INFILE_CREATED.replace("Z", "+00:00")).replace(tzinfo=None)
    cp.modified = datetime.fromisoformat(INFILE_MODIFIED.replace("Z", "+00:00")).replace(
        tzinfo=None
    )
    doc.save(p)
    files["docx"] = p

    from pypdf import PdfWriter

    p = src / "paper.pdf"
    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    w.add_metadata({"/CreationDate": "D:20230110090000Z", "/ModDate": "D:20250314043000Z"})
    with open(p, "wb") as fh:
        w.write(fh)
    files["pdf"] = p

    from PIL import Image

    p = src / "photo.jpg"
    im = Image.new("RGB", (16, 16), (120, 80, 40))
    exif = Image.Exif()
    exif.get_ifd(0x8769)[36867] = EXIF_ORIGINAL
    im.save(p, exif=exif)
    files["jpg"] = p

    p = src / "bundle.zip"
    with zipfile.ZipFile(p, "w") as zf:
        info = zipfile.ZipInfo("inner.txt", date_time=ZIP_ENTRY)
        zf.writestr(info, "inner")
    files["zip"] = p

    for f in files.values():
        set_file_times(f, created=SRC_CREATED, modified=SRC_MODIFIED, accessed=SRC_ACCESSED)
    return files


# ---------------------------------------------------------------------------
# Tools
# ---------------------------------------------------------------------------


def _ps(command: str) -> None:
    subprocess.run(
        ["powershell", "-NoProfile", "-NonInteractive", "-Command", command],
        check=True,
        capture_output=True,
        text=True,
    )


def _q(path: Path) -> str:
    return "'" + str(path).replace("'", "''") + "'"


def tool_copy_item(src: Path, dst_dir: Path) -> Path:
    _ps(f"Copy-Item -LiteralPath {_q(src)} -Destination {_q(dst_dir)}")
    return dst_dir / src.name


def tool_cmd_copy(src: Path, dst_dir: Path) -> Path:
    subprocess.run(
        ["cmd", "/c", "copy", "/Y", str(src), str(dst_dir)], check=True, capture_output=True
    )
    return dst_dir / src.name


def tool_xcopy(src: Path, dst_dir: Path) -> Path:
    subprocess.run(
        ["xcopy", str(src), str(dst_dir) + "\\", "/Y", "/Q"], check=True, capture_output=True
    )
    return dst_dir / src.name


def _robocopy(src: Path, dst_dir: Path, *flags: str) -> Path:
    proc = subprocess.run(
        [
            "robocopy",
            str(src.parent),
            str(dst_dir),
            src.name,
            "/NJH",
            "/NJS",
            "/NFL",
            "/NDL",
            *flags,
        ],
        capture_output=True,
    )
    if proc.returncode >= 8:
        raise RuntimeError(f"robocopy failed: {proc.returncode}")
    return dst_dir / src.name


def tool_robocopy_default(src: Path, dst_dir: Path) -> Path:
    return _robocopy(src, dst_dir)


def tool_robocopy_copy_d(src: Path, dst_dir: Path) -> Path:
    return _robocopy(src, dst_dir, "/COPY:D")  # data only


def tool_shell_copyhere(src: Path, dst_dir: Path) -> Path:
    """Explorer's own copy engine via Shell.Application (a scripted proxy for drag-copy)."""
    _ps(
        "$sh = New-Object -ComObject Shell.Application; "
        f"$sh.NameSpace({_q(dst_dir)}).CopyHere({_q(src)}, 16+4+1024); "
        f"$t=0; while(-not (Test-Path -LiteralPath {_q(dst_dir / src.name)}) -and $t -lt 100) {{ Start-Sleep -Milliseconds 100; $t++ }}"
    )
    time.sleep(0.5)
    return dst_dir / src.name


def tool_shell_movehere(src: Path, dst_dir: Path) -> Path:
    """Explorer move within the same volume (a rename)."""
    tmp = src.parent / ("mv_" + src.name)
    shutil.copy2(src, tmp)
    set_file_times(tmp, created=SRC_CREATED, modified=SRC_MODIFIED, accessed=SRC_ACCESSED)
    _ps(
        "$sh = New-Object -ComObject Shell.Application; "
        f"$sh.NameSpace({_q(dst_dir)}).MoveHere({_q(tmp)}, 16+4+1024); "
        f"$t=0; while(-not (Test-Path -LiteralPath {_q(dst_dir / tmp.name)}) -and $t -lt 100) {{ Start-Sleep -Milliseconds 100; $t++ }}"
    )
    time.sleep(0.5)
    return dst_dir / tmp.name


def tool_shutil_copy(src: Path, dst_dir: Path) -> Path:
    return Path(shutil.copy(src, dst_dir))


def tool_shutil_copy2(src: Path, dst_dir: Path) -> Path:
    return Path(shutil.copy2(src, dst_dir))


def tool_shutil_copyfile(src: Path, dst_dir: Path) -> Path:
    dst = dst_dir / src.name
    shutil.copyfile(src, dst)
    return dst


def tool_compress_expand_archive(src: Path, dst_dir: Path) -> Path:
    archive = dst_dir / (src.stem + "_ps.zip")
    _ps(f"Compress-Archive -LiteralPath {_q(src)} -DestinationPath {_q(archive)} -Force")
    out = dst_dir / "expanded"
    _ps(f"Expand-Archive -LiteralPath {_q(archive)} -DestinationPath {_q(out)} -Force")
    return out / src.name


def tool_python_zip_extract(src: Path, dst_dir: Path) -> Path:
    archive = dst_dir / (src.stem + "_py.zip")
    with zipfile.ZipFile(archive, "w") as zf:
        zf.write(src, src.name)
    out = dst_dir / "pyextract"
    with zipfile.ZipFile(archive) as zf:
        zf.extractall(out)  # Python's extractall does NOT restore entry times
    return out / src.name


def tool_tar_zip_roundtrip(src: Path, dst_dir: Path) -> Path:
    archive = dst_dir / (src.stem + "_tar.zip")
    subprocess.run(
        ["tar.exe", "-a", "-cf", str(archive), "-C", str(src.parent), src.name],
        check=True,
        capture_output=True,
    )
    out = dst_dir / "tarextract"
    out.mkdir(exist_ok=True)
    subprocess.run(
        ["tar.exe", "-xf", str(archive), "-C", str(out)], check=True, capture_output=True
    )
    return out / src.name


def tool_git_clone(src: Path, dst_dir: Path) -> Path:
    repo = dst_dir / "repo"
    repo.mkdir()
    shutil.copy2(src, repo / src.name)
    set_file_times(
        repo / src.name, created=SRC_CREATED, modified=SRC_MODIFIED, accessed=SRC_ACCESSED
    )
    git = ["git", "-c", "safe.directory=*", "-c", "user.name=t", "-c", "user.email=t@t"]
    subprocess.run([*git, "init", "-q"], cwd=repo, check=True, capture_output=True)
    subprocess.run([*git, "add", "."], cwd=repo, check=True, capture_output=True)
    subprocess.run([*git, "commit", "-qm", "x"], cwd=repo, check=True, capture_output=True)
    clone = dst_dir / "clone"
    subprocess.run([*git, "clone", "-q", str(repo), str(clone)], check=True, capture_output=True)
    return clone / src.name


def tool_safe_save_tunnel(src: Path, dst_dir: Path) -> Path:
    """Editor-style safe save: write temp, delete original, rename temp (NTFS tunnelling)."""
    target = dst_dir / src.name
    shutil.copy2(src, target)
    set_file_times(target, created=SRC_CREATED, modified=SRC_MODIFIED, accessed=SRC_ACCESSED)
    tmp = dst_dir / (src.name + ".tmp")
    tmp.write_bytes(target.read_bytes())
    target.unlink()
    tmp.rename(target)
    return target


def tool_in_place_append(src: Path, dst_dir: Path) -> Path:
    """An ordinary edit: open, append, close (txt only — would corrupt binaries)."""
    target = dst_dir / src.name
    shutil.copy2(src, target)
    set_file_times(target, created=SRC_CREATED, modified=SRC_MODIFIED, accessed=SRC_ACCESSED)
    if target.suffix == ".txt":
        with open(target, "a", encoding="utf-8") as fh:
            fh.write("edit\n")
    else:
        with open(target, "r+b") as fh:  # rewrite identical bytes: an edit that changes nothing
            data = fh.read()
            fh.seek(0)
            fh.write(data)
    return target


def tool_python_docx_resave(src: Path, dst_dir: Path) -> Path | None:
    """Open and save with python-docx (a library 'Save' that does not touch core.xml)."""
    if src.suffix != ".docx":
        return None
    from docx import Document

    target = dst_dir / ("resaved_" + src.name)
    Document(str(src)).save(str(target))
    return target


TOOLS = {
    "copy_item_ps": tool_copy_item,
    "cmd_copy": tool_cmd_copy,
    "xcopy": tool_xcopy,
    "robocopy_default_DAT": tool_robocopy_default,
    "robocopy_COPY_D": tool_robocopy_copy_d,
    "shell_copyhere_explorer_engine": tool_shell_copyhere,
    "shell_movehere_same_volume": tool_shell_movehere,
    "shutil_copy": tool_shutil_copy,
    "shutil_copy2": tool_shutil_copy2,
    "shutil_copyfile": tool_shutil_copyfile,
    "compress_expand_archive_ps": tool_compress_expand_archive,
    "python_zipfile_extractall": tool_python_zip_extract,
    "tar_zip_roundtrip": tool_tar_zip_roundtrip,
    "git_clone": tool_git_clone,
    "safe_save_ntfs_tunnel": tool_safe_save_tunnel,
    "in_place_edit": tool_in_place_append,
    "python_docx_resave": tool_python_docx_resave,
}


# ---------------------------------------------------------------------------
# Classification
# ---------------------------------------------------------------------------


def _classify(
    observed: datetime, source: datetime, run_start: datetime, source_now: datetime | None = None
) -> str:
    if abs(observed - source) <= timedelta(seconds=2):
        return "kept"
    if source_now is not None and abs(observed - source_now) <= timedelta(seconds=2):
        return "source_current"  # copied the source's value as it was at copy time
    if (
        observed >= run_start - timedelta(seconds=2)
        and observed <= datetime.now(timezone.utc) + NOW_TOLERANCE
    ):
        return "now"
    return f"other({observed.isoformat()})"


def _has_motw(path: Path) -> bool:
    try:
        with open(str(path) + ":Zone.Identifier", encoding="utf-8"):
            return True
    except OSError:
        return False


def measure(work: Path) -> dict:
    if work.exists():
        import os
        import stat

        def _force(func, path, _exc):  # git marks object files read-only
            os.chmod(path, stat.S_IWRITE)
            func(path)

        shutil.rmtree(work, onexc=_force)
    fixtures = make_fixtures(work / "src")
    src_fs = {fmt: get_file_times(p) for fmt, p in fixtures.items()}
    src_infile = {fmt: read_infile_times(p) for fmt, p in fixtures.items()}

    rows = []
    for tool_name, tool in TOOLS.items():
        for fmt, src in fixtures.items():
            dst_dir = work / "dst" / tool_name / fmt
            dst_dir.mkdir(parents=True)
            src_before = get_file_times(src)
            run_start = datetime.now(timezone.utc)
            time.sleep(1.1)  # make "now" distinguishable from a same-second source
            try:
                dst = tool(src, dst_dir)
            except Exception as exc:  # noqa: BLE001 - record, keep measuring
                rows.append({"tool": tool_name, "format": fmt, "error": repr(exc)})
                continue
            if dst is None:
                continue
            fs = get_file_times(dst)
            s = src_fs[fmt]
            infile_after = read_infile_times(dst)
            rows.append(
                {
                    "tool": tool_name,
                    "format": fmt,
                    "birth": _classify(fs.created, s.created, run_start),
                    "mtime": _classify(fs.modified, s.modified, run_start),
                    "atime": _classify(fs.accessed, s.accessed, run_start, src_before.accessed),
                    "birth_gt_mtime": fs.created > fs.modified,
                    "infile": {
                        key: (
                            "kept"
                            if infile_after.get(key) == value
                            else f"changed({infile_after.get(key)})"
                        )
                        for key, value in src_infile[fmt].items()
                    },
                    "motw": _has_motw(dst),
                }
            )
    return {
        "schema": "contextfs.transfer_table/1",
        "measured_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "machine": {
            "os": platform.platform(),
            "volume_fs": "NTFS (D:)",
            "python": platform.python_version(),
        },
        "source_times_local": {
            "created": SRC_CREATED.isoformat(),
            "modified": SRC_MODIFIED.isoformat(),
            "accessed": SRC_ACCESSED.isoformat(),
        },
        "source_infile": src_infile,
        "rows": rows,
        "pending_user": [{"tool": t, "procedure": d} for t, d in PENDING_USER],
        "not_covered": [
            "cross-volume moves (would need a second writable volume; C: is excluded by project rule)",
            "network shares / SMB",
        ],
    }


def to_markdown(table: dict) -> str:
    fmts = FORMATS
    lines = [
        f"# Transfer table T — measured {table['measured_at']}",
        "",
        f"Machine: {table['machine']['os']}, {table['machine']['volume_fs']}. "
        "Source FS times: created < modified < accessed, all in 2025. "
        "`kept` = unchanged, `now` = set to transfer time, "
        "`source_current` = copied the source's value at copy time (source atime moves when tools read it).",
        "",
        "| tool | birth | mtime | atime | birth>mtime | in-file times | MotW |",
        "|---|---|---|---|---|---|---|",
    ]
    by_tool: dict[str, list[dict]] = {}
    for row in table["rows"]:
        by_tool.setdefault(row["tool"], []).append(row)
    for tool, rows in by_tool.items():
        ok = [r for r in rows if "error" not in r]
        if not ok:
            lines.append(f"| {tool} | error: {rows[0]['error'][:60]} | | | | | |")
            continue

        def agg(key, ok=ok):
            vals = sorted({str(r[key]) for r in ok})
            return "/".join(vals)

        infile = sorted({f"{k}:{v}" for r in ok for k, v in r["infile"].items() if v != "kept"})
        lines.append(
            f"| {tool} | {agg('birth')} | {agg('mtime')} | {agg('atime')} | {agg('birth_gt_mtime')} | "
            f"{'all kept' if not infile else '; '.join(infile)} | {agg('motw')} |"
        )
    lines += [
        "",
        "Formats per tool: " + ", ".join(fmts) + ".",
        "",
        "## Pending (needs the user)",
        "",
    ]
    lines += [f"- `{p['tool']}`: {p['procedure']}" for p in table["pending_user"]]
    lines += ["", "## Not covered", ""] + [f"- {n}" for n in table["not_covered"]]
    return "\n".join(lines) + "\n"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--work", type=Path, default=PROJECT_ROOT / ".tools" / "exp" / "transfer")
    parser.add_argument("--out", type=Path, default=PROJECT_ROOT / "data" / "transfer_table")
    args = parser.parse_args()

    table = measure(args.work)
    args.out.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(table, indent=2, sort_keys=True)
    (args.out / "T_v1.json").write_text(payload + "\n", encoding="utf-8")
    (args.out / "T_v1.md").write_text(to_markdown(table), encoding="utf-8")
    digest = hashlib.sha256((payload + "\n").encode("utf-8")).hexdigest()
    (args.out / "T_v1.sha256").write_text(f"{digest}  T_v1.json\n", encoding="utf-8")
    print(to_markdown(table))
    print(f"sha256 {digest}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
